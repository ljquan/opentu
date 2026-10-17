from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\ZhuanZ1\Desktop\WorkBuddy配置教程_完善版.docx")
OUTPUT = Path(r"C:\Users\ZhuanZ1\Desktop\WorkBuddy配置教程_完善图文版.docx")
MEDIA = Path(r"C:\Users\ZhuanZ1\AppData\Local\Temp\codex_workbuddy_source_media_20260829\word\media")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, name="Microsoft YaHei", size=10.5, color=None, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def insert_after(paragraph, new_paragraph):
    paragraph._p.addnext(new_paragraph._p)


def add_figure_after(doc, anchor, filename, caption, width=6.25):
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(4)
    figure.paragraph_format.space_after = Pt(2)
    run = figure.add_run()
    run.add_picture(str(MEDIA / filename), width=Inches(width))
    insert_after(anchor, figure)

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, color="666666")
    insert_after(figure, caption_paragraph)
    return caption_paragraph


def add_note_after(doc, anchor, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=9.5, color="7A4E00")
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "E6B566")
    pBdr.append(bottom)
    pPr.append(pBdr)
    insert_after(anchor, p)
    return p


doc = Document(str(SOURCE))

# Locate the existing numbered instructions so the figures remain tied to the
# relevant action even if the surrounding prose is revised later.
paras = doc.paragraphs
anchors = {}
for p in paras:
    text = p.text.strip()
    for prefix, key in [
        ("1. 打开 WorkBuddy", "open"),
        ("2. 选择“自定义配置”", "custom"),
        ("3. 填写接口地址", "fields"),
        ("4. 填写 API Key", "key"),
        ("5. 选择模型", "model"),
        ("6. 点击“测试连接”", "test"),
        ("7. 测试成功后保存配置", "save"),
    ]:
        if text.startswith(prefix):
            anchors[key] = p

if len(anchors) != 7:
    raise RuntimeError(f"Expected 7 tutorial anchors, found {sorted(anchors)}")

# Put the first two navigation screenshots directly after their actions.
add_figure_after(doc, anchors["open"], "image1.png", "图 1  在模型选择菜单中打开“配置自定义模型”")
add_figure_after(doc, anchors["custom"], "image2.png", "图 2  在“添加模型”窗口中选择“自定义 / Custom”")

# The field form is shared by the address, key, and model steps, so show it
# once after the model step rather than repeating the same large screenshot.
field_caption = (
    "图 3  自定义模型字段填写示例（截图为界面示意；接口地址请以本文正文和服务商当前提供的地址为准）"
)
add_figure_after(doc, anchors["model"], "image3.png", field_caption)

# Place the validation and post-save screenshots alongside the final actions.
add_figure_after(doc, anchors["test"], "image4.png", "图 4  点击“测试连接”，出现“连接成功”后保存配置")
last_save_figure = add_figure_after(doc, anchors["save"], "image5.png", "图 5  保存后，在模型列表中选中刚刚配置的模型")
add_figure_after(doc, last_save_figure, "image6.png", "图 6  新建对话并确认底部显示已选中的自定义模型")

# Add a short, visible clarification after the address bullets. This prevents
# the older address visible in image3 from being copied by readers.
address_anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("• 如果后台已经给了固定地址"))
add_note_after(
    doc,
    address_anchor,
    "截图提示：部分旧版本界面截图可能显示不完整的旧地址。请优先使用本文列出的完整地址，尤其确认末尾包含 /v1。",
)

# The final line was a drafting note, not publishable tutorial content.
for p in list(doc.paragraphs):
    if p.text.strip().startswith("建议：如果你希望这份教程更适合发布"):
        p._element.getparent().remove(p._element)

# Mark the comparison table's first row as a repeating header for Word and
# screen-reader navigation.
for table in doc.tables:
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))

# Add descriptive alt text to every newly inserted picture for accessibility.
alt_texts = {
    "image1.png": "WorkBuddy 模型菜单，箭头指向配置自定义模型",
    "image2.png": "WorkBuddy 添加模型窗口，箭头指向自定义 Custom",
    "image3.png": "WorkBuddy 自定义模型编辑窗口，展示接口地址、API Key 和模型名称字段",
    "image4.png": "WorkBuddy 自定义模型编辑窗口，显示连接成功提示、测试连接和保存按钮",
    "image5.png": "WorkBuddy 模型选择菜单，显示已保存的自定义模型",
    "image6.png": "WorkBuddy 对话界面底部显示当前选中的自定义模型",
}
for shape, description in zip(doc.inline_shapes, alt_texts.values()):
    docPr = shape._inline.docPr
    docPr.set("descr", description)

doc.core_properties.title = "WorkBuddy 配置教程（图文版）"
doc.core_properties.subject = "WorkBuddy 自定义 API 配置、测试与模型选择"
doc.core_properties.author = ""
doc.core_properties.comments = ""
doc.save(str(OUTPUT))
print(OUTPUT)
